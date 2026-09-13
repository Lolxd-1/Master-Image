"""
T-1.5 independent verification.

The JS harness checks the writer with code that shares assumptions with the writer. This pass
re-reads the generated files with openpyxl and Python's XML parser -- implementations that share
nothing with our engine -- so a bug in our XML handling cannot hide behind itself.
"""
import sys, io, os, zipfile, warnings
import xml.etree.ElementTree as ET
import openpyxl

sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8', errors='replace')

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
MASTER = os.path.join(ROOT, 'Master excel.xlsx')
OUT = os.path.join(ROOT, '.test-output')

passed = failed = 0
def check(cid, desc, fn):
    global passed, failed
    try:
        note = fn()
        print(f"  PASS  {cid}  {desc}" + (f"  -- {note}" if note else ""))
        passed += 1
    except Exception as e:
        print(f"  FAIL  {cid}  {desc}\n          {type(e).__name__}: {e}")
        failed += 1

def rows_of(path, sheet='bulk_upload_template'):
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = wb[sheet]
    out = []
    for r in ws.iter_rows(min_row=2, values_only=True):
        if any(c is not None and str(c).strip() != "" for c in r):
            out.append(r)
    wb.close()
    return out

print("\nT-1.5  Independent verification (openpyxl + ElementTree)\n")

master_rows = rows_of(MASTER)

# ---------------------------------------------------------------- loads cleanly

def t_loads():
    with warnings.catch_warnings(record=True) as w:
        warnings.simplefilter("always")
        wb = openpyxl.load_workbook(os.path.join(OUT, 'all.xlsx'))
        names = wb.sheetnames
        hidden = wb['DataSheet'].sheet_state
        wb.close()
    msgs = [str(x.message) for x in w]
    assert not msgs, f"openpyxl emitted warnings: {msgs}"
    assert names == ['important_instructions', 'bulk_upload_template', 'DataSheet'], names
    assert hidden == 'hidden', f"DataSheet should stay hidden, got {hidden}"
    return f"3 sheets, DataSheet hidden, no warnings"
check('T-1.5a', 'output loads with no warnings and keeps its sheet structure', t_loads)

# ---------------------------------------------------------------- every XML part is well-formed

def t_wellformed():
    z = zipfile.ZipFile(os.path.join(OUT, 'all.xlsx'))
    bad = z.testzip()
    assert bad is None, f"corrupt zip entry: {bad}"
    n = 0
    for name in z.namelist():
        if name.endswith('.xml') or name.endswith('.rels'):
            ET.fromstring(z.read(name))   # raises on malformed XML
            n += 1
    return f"{n} XML parts parse cleanly, zip CRCs valid"
check('T-1.5b', 'every XML part is well-formed and the zip is not corrupt', t_wellformed)

# ---------------------------------------------------------------- full round-trip equality

def t_full():
    out = rows_of(os.path.join(OUT, 'all.xlsx'))
    assert len(out) == len(master_rows) == 650, f"{len(out)} vs {len(master_rows)}"
    for i, (a, b) in enumerate(zip(out, master_rows)):
        assert len(a) == len(b) == 25, f"row {i+2} has {len(a)}/{len(b)} columns, expected 25"
        for j in range(25):
            if a[j] != b[j]:
                raise AssertionError(
                    f"row {i+2} col {openpyxl.utils.get_column_letter(j+1)}: "
                    f"{a[j]!r} != {b[j]!r}")
    return "650 rows x 25 columns identical to source"
check('T-1.5c', 'all-selected export equals the master cell for cell', t_full)

# ---------------------------------------------------------------- types preserved

def t_types():
    out = rows_of(os.path.join(OUT, 'all.xlsx'))
    for i, r in enumerate(out):
        for j, col in ((4, 'E/MRP'), (5, 'F/SellingPrice')):
            assert isinstance(r[j], (int, float)), \
                f"row {i+2} {col} is {type(r[j]).__name__} ({r[j]!r}) -- SmartBiz rejects text prices"
        assert isinstance(r[0], str) and len(r[0]) == 36, f"row {i+2} SKU malformed: {r[0]!r}"
    return "650 rows: prices numeric, SKUs intact"
check('T-1.5d', 'numeric columns stay numeric, SKU IDs stay intact', t_types)

# ---------------------------------------------------------------- validations readable by openpyxl

def t_validations():
    wb = openpyxl.load_workbook(os.path.join(OUT, 'all.xlsx'))
    ws = wb['bulk_upload_template']
    dvs = ws.data_validations.dataValidation
    wb.close()
    assert len(dvs) == 23, f"expected 23 data validations, found {len(dvs)}"
    formulas = {str(d.formula1) for d in dvs if d.formula1}
    for named in ('BCat', 'Best_Option'):
        assert named in formulas, f"named-range dropdown {named} missing (found {sorted(formulas)[:8]})"
    ranges = {str(d.sqref) for d in dvs}
    assert any('G2:G25000' in r for r in ranges), "business-category dropdown range missing"
    return f"{len(dvs)} validations, BCat + Best_Option dropdowns live"
check('T-1.5e', 'dropdowns are readable as dropdowns, not just present as text', t_validations)

# ---------------------------------------------------------------- subset correctness

def t_subset():
    out = rows_of(os.path.join(OUT, 'subset.xlsx'))
    expected = master_rows[::5]
    assert len(out) == len(expected), f"{len(out)} rows vs expected {len(expected)}"
    for i, (a, b) in enumerate(zip(out, expected)):
        assert a[0] == b[0], f"row {i+2} SKU {a[0]} != {b[0]}"
        assert a[3] == b[3], f"row {i+2} name mismatch"
        assert a[4] == b[4] and a[5] == b[5], f"row {i+2} price mismatch"
    return f"{len(out)} rows, SKUs and prices match source order"
check('T-1.5f', 'subset export contains exactly the chosen rows in order', t_subset)

# ---------------------------------------------------------------- image URLs untouched

def t_images():
    out = rows_of(os.path.join(OUT, 'all.xlsx'))
    n = 0
    for i, (a, b) in enumerate(zip(out, master_rows)):
        assert a[15] == b[15], f"row {i+2} image URL changed:\n  {a[15]}\n  {b[15]}"
        if a[15]:
            assert '_SX' not in a[15] and '_SL' not in a[15], f"resized URL leaked: {a[15]}"
            n += 1
    return f"{n} URLs byte-identical, none resized"
check('T-1.5g', 'image URLs reach the file exactly as they came in', t_images)

# ---------------------------------------------------------------- special characters

def t_special():
    out = rows_of(os.path.join(OUT, 'special.xlsx'))
    names = [r[3] for r in out]
    lays = [n for n in names if "LAY'S" in n]
    assert lays, f"apostrophe product missing; sample: {names[:3]}"
    src = {r[3] for r in master_rows}
    for n in names:
        assert n in src, f"name altered in export: {n!r}"
    return f"{len(out)} rows, e.g. {lays[0]}"
check('T-1.5h', 'apostrophes and ampersands survive the round trip', t_special)

# ---------------------------------------------------------------- overrides (T-1.14/T-1.20)

def t_override_mixed():
    path = os.path.join(OUT, 'override-mixed.xlsx')
    assert os.path.exists(path), "override-mixed.xlsx missing -- run verify-xlsx.mjs first"
    with warnings.catch_warnings(record=True) as w:
        warnings.simplefilter("always")
        wb = openpyxl.load_workbook(path)
        ws = wb['bulk_upload_template']
        dvs = ws.data_validations.dataValidation
        # edited row reads back with exact values and numeric types
        rows = list(ws.iter_rows(min_row=2, values_only=True))
        rows = [r for r in rows if any(c is not None and str(c).strip() != "" for c in r)]
        wb.close()
    msgs = [str(x.message) for x in w]
    assert not msgs, f"openpyxl warnings on overrides file: {msgs}"
    assert len(dvs) == 23, f"expected 23 validations, found {len(dvs)}"
    names = [r[3] for r in rows]
    assert 'EDITED' in names, f"edited name missing: {names[:3]}"
    idx = names.index('EDITED')
    assert isinstance(rows[idx][4], (int, float)) and rows[idx][4] == 99, f"MRP wrong: {rows[idx][4]!r}"
    assert isinstance(rows[idx][5], (int, float)) and rows[idx][5] == 88, f"price wrong: {rows[idx][5]!r}"
    return f"{len(rows)} rows, EDITED reads back MRP=99 price=88 numeric"
check('T-1.20b', 'overrides file loads clean and edited values read back exactly', t_override_mixed)

def t_override_price():
    path = os.path.join(OUT, 'override-price.xlsx')
    assert os.path.exists(path), "override-price.xlsx missing -- run verify-xlsx.mjs first"
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = wb['bulk_upload_template']
    rows = [r for r in ws.iter_rows(min_row=2, values_only=True)
            if any(c is not None and str(c).strip() != "" for c in r)]
    wb.close()
    assert len(rows) == 1, f"expected 1 row, got {len(rows)}"
    assert rows[0][5] == 38 and isinstance(rows[0][5], (int, float)), f"price wrong: {rows[0][5]!r}"
    return "price=38 numeric"
check('T-1.13b', 'overridden price reads back as a number via openpyxl', t_override_price)

print(f"\n  {passed} passed, {failed} failed\n")
sys.exit(1 if failed else 0)
